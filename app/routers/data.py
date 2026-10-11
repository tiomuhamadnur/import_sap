from __future__ import annotations

import csv
import tempfile
import uuid
from datetime import date, datetime
from decimal import Decimal
from math import ceil
from pathlib import Path
from typing import Annotated
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Form, HTTPException, Request
from fastapi.responses import FileResponse, RedirectResponse
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import delete, select
from starlette.background import BackgroundTask

from app.dependencies import AdminDep, CsrfDep, DbDep, RequireUserDep
from app.models import AuditLog
from app.services.dataviews import (
    EntityView,
    build_query,
    count_query,
    editable_columns,
    existing_values,
    get_view,
)
from app.templating import render

router = APIRouter(tags=["data"])

PER_PAGE_CHOICES = (25, 50, 100)

CSV_MEDIA = "text/csv"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _make_qs(params: dict[str, str]):
    def qs(**overrides):
        merged = {k: v for k, v in params.items() if k not in overrides}
        merged.update({k: v for k, v in overrides.items() if v not in (None, "")})
        return "?" + urlencode(merged)

    return qs


def get_view_or_404(slug: str) -> EntityView:
    view = get_view(slug)
    if view is None:
        raise HTTPException(status_code=404, detail="Unknown data page")
    return view


ViewDep = Annotated[EntityView, Depends(get_view_or_404)]


def _get_record(db, view: EntityView, public_id: uuid.UUID):
    return db.scalar(select(view.model).where(view.model.public_id == public_id))


def _filters(request: Request) -> dict[str, str]:
    return {
        key[2:]: value
        for key, value in request.query_params.items()
        if key.startswith("f_") and value.strip()
    }


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, Decimal):
        return str(value)
    return str(value)


def _relations_for_rows(db, view: EntityView, rows: list) -> dict[str, set[str]]:
    out: dict[str, set[str]] = {}
    for relation in view.relations:
        values = {str(getattr(r, relation.field)) for r in rows if getattr(r, relation.field)}
        out[relation.field] = existing_values(db, relation, values)
    return out


@router.get("/{slug}")
def data_list(
    request: Request,
    db: DbDep,
    view: ViewDep,
    current_user: RequireUserDep,
    q: str | None = None,
    sort: str | None = None,
    dir: str = "asc",
    page: int = 1,
    per_page: int = 25,
):
    filters = _filters(request)
    per_page = per_page if per_page in PER_PAGE_CHOICES else 25
    page = max(1, page)
    direction = "desc" if dir == "desc" else "asc"

    stmt = build_query(view, filters, q, sort, direction)
    total = count_query(db, stmt)
    pages = max(1, ceil(total / per_page))
    page = min(page, pages)
    rows = list(db.scalars(stmt.offset((page - 1) * per_page).limit(per_page)).all())

    return render(
        request,
        "data_list.html",
        active=view.title,
        current_user=current_user,
        view=view,
        rows=rows,
        related=_relations_for_rows(db, view, rows),
        relations_by_field={r.field: r for r in view.relations},
        filters=filters,
        q=q or "",
        sort=sort or view.key_field,
        dir=direction,
        page=page,
        pages=pages,
        per_page=per_page,
        per_page_choices=PER_PAGE_CHOICES,
        total=total,
        qs=_make_qs(dict(request.query_params)),
    )


@router.get("/{slug}/export")
def data_export(
    request: Request,
    db: DbDep,
    view: ViewDep,
    current_user: RequireUserDep,
    format: str = "csv",
    q: str | None = None,
    sort: str | None = None,
    dir: str = "asc",
):
    filters = _filters(request)
    direction = "desc" if dir == "desc" else "asc"
    stmt = build_query(view, filters, q, sort, direction)
    fields = [c.field for c in view.columns]
    headers = [c.label for c in view.columns]
    # Only fetch the exported columns (skips heavy JSONB raw_data, avoids ORM overhead).
    row_stmt = stmt.with_only_columns(*[getattr(view.model, f) for f in fields])

    if format == "xlsx":
        return _export_xlsx(db, view, row_stmt, headers, fields)
    return _export_csv(db, view, row_stmt, headers, fields)


def _export_rows(db, stmt):
    for row in db.execute(stmt.execution_options(yield_per=1000)):
        yield tuple(row)


def _export_csv(db, view: EntityView, stmt, headers, fields):
    tmp = tempfile.NamedTemporaryFile(suffix=".csv", delete=False)
    tmp.close()
    with open(tmp.name, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(headers)
        for row in _export_rows(db, stmt):
            writer.writerow([_cell(v) for v in row])
    return FileResponse(
        tmp.name,
        media_type=f"{CSV_MEDIA}; charset=utf-8",
        filename=f"{view.slug}.csv",
        background=BackgroundTask(lambda p=tmp.name: Path(p).unlink(missing_ok=True)),
    )


def _xlsx_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, int | float):
        return value
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, date):
        return value
    return str(value)


def _export_xlsx(db, view: EntityView, stmt, headers, fields):
    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=view.title[:31])
    header_cells = [WriteOnlyCell(ws, value=h) for h in headers]
    for cell in header_cells:
        cell.font = Font(bold=True)
    ws.append(header_cells)
    n = 1
    for row in _export_rows(db, stmt):
        ws.append([_xlsx_value(v) for v in row])
        n += 1
    ws.freeze_panes = "A2"
    if n > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n}"
    wb.save(tmp.name)
    return FileResponse(
        tmp.name,
        media_type=XLSX_MEDIA,
        filename=f"{view.slug}.xlsx",
        background=BackgroundTask(lambda p=tmp.name: Path(p).unlink(missing_ok=True)),
    )


@router.get("/{slug}/bulk-delete")
def bulk_delete_get():  # pragma: no cover - guard, real action is POST
    raise HTTPException(status_code=405, detail="Use POST")


@router.post("/{slug}/bulk-delete")
def bulk_delete(
    db: DbDep,
    view: ViewDep,
    current_user: AdminDep,
    _csrf: CsrfDep,
    ids: Annotated[list[int] | None, Form()] = None,
):
    ids = ids or []
    deleted = 0
    if ids:
        result = db.execute(delete(view.model).where(view.model.id.in_(ids)))
        deleted = result.rowcount or 0
        db.add(
            AuditLog(
                user_id=current_user.id,
                action=f"{view.data_type}.bulk_delete",
                object_type=view.model.__tablename__,
                result="success",
                detail={"count": deleted, "ids": ids[:100]},
            )
        )
        db.commit()
    return RedirectResponse(f"/{view.slug}?deleted={deleted}", status_code=303)


@router.get("/{slug}/{public_id}")
def data_detail(
    request: Request,
    db: DbDep,
    view: ViewDep,
    current_user: RequireUserDep,
    public_id: uuid.UUID,
):
    record = _get_record(db, view, public_id)
    if record is None:
        raise HTTPException(status_code=404, detail="Record not found")

    relations = []
    for relation in view.relations:
        value = getattr(record, relation.field)
        exists = False
        if value:
            exists = bool(existing_values(db, relation, {str(value)}))
        relations.append(
            {
                "relation": relation,
                "value": value,
                "exists": exists,
                "href": f"/{relation.slug}?f_{relation.target_field}={value}",
            }
        )

    fields_meta = [
        {"field": c.field, "label": c.label, "value": getattr(record, c.field)}
        for c in view.columns
    ]
    return render(
        request,
        "data_detail.html",
        active=view.title,
        current_user=current_user,
        view=view,
        record=record,
        relations=relations,
        fields_meta=fields_meta,
        error=request.query_params.get("error"),
    )


@router.get("/{slug}/{public_id}/edit")
def data_edit(
    request: Request,
    db: DbDep,
    view: ViewDep,
    current_user: AdminDep,
    public_id: uuid.UUID,
):
    record = _get_record(db, view, public_id)
    if record is None:
        raise HTTPException(status_code=404, detail="Record not found")
    editable = editable_columns(view)
    is_modal = request.headers.get("HX-Request") == "true"
    return render(
        request,
        "data_edit_modal.html" if is_modal else "data_edit.html",
        active=view.title,
        current_user=current_user,
        view=view,
        record=record,
        editable=editable,
        error=None,
    )


@router.post("/{slug}/{public_id}")
async def data_update(
    request: Request,
    db: DbDep,
    view: ViewDep,
    current_user: AdminDep,
    _csrf: CsrfDep,
    public_id: uuid.UUID,
):
    record = _get_record(db, view, public_id)
    if record is None:
        raise HTTPException(status_code=404, detail="Record not found")

    form = await request.form()
    changes = {}
    for field_name in view.editable:
        new = (form.get(field_name) or "").strip()
        old = getattr(record, field_name)
        old_str = "" if old is None else str(old)
        if new != old_str:
            setattr(record, field_name, new or None)
            changes[field_name] = {"from": old_str, "to": new}

    if changes:
        db.add(
            AuditLog(
                user_id=current_user.id,
                action=f"{view.data_type}.update",
                object_type=view.model.__tablename__,
                object_id=str(record.public_id),
                result="success",
                detail={"changes": changes},
            )
        )
        db.commit()
    return RedirectResponse(f"/{view.slug}/{record.public_id}", status_code=303)


@router.post("/{slug}/{public_id}/delete")
def data_delete(
    db: DbDep,
    view: ViewDep,
    current_user: AdminDep,
    _csrf: CsrfDep,
    public_id: uuid.UUID,
):
    record = _get_record(db, view, public_id)
    if record is None:
        raise HTTPException(status_code=404, detail="Record not found")
    key = str(getattr(record, view.key_field))
    db.delete(record)
    db.add(
        AuditLog(
            user_id=current_user.id,
            action=f"{view.data_type}.delete",
            object_type=view.model.__tablename__,
            object_id=str(record.public_id),
            result="success",
            detail={"key": key},
        )
    )
    db.commit()
    return RedirectResponse(f"/{view.slug}", status_code=303)
