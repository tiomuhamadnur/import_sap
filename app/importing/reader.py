from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook

DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y")
_NO_VALUE = ("", "0", "00:00:00", "0,00", "0,000", "000000")


def sanitize(value):
    """Convert a raw cell value into a JSON-serializable form for raw_data storage."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return value


def is_empty(v) -> bool:
    if v is None:
        return True
    if isinstance(v, int | float) and v in (0, 0.0):
        return True
    s = str(v).strip()
    return s == "" or (re.fullmatch(r"0+(,\d+)?", s) is not None and s.replace(",", ".") == "0")


def parse_date(value) -> datetime | None:
    from datetime import date

    if value is None:
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    s = str(value).strip()
    if s in _NO_VALUE:
        return None
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def parse_number(value) -> float | None:
    if value is None:
        return None
    if isinstance(value, int | float) and not isinstance(value, bool):
        return float(value)
    s = str(value).strip().replace("Rp ", "").replace("IDR ", "").strip()
    if s in ("",):
        return None
    s = re.sub(r"[^\d,.\-+]", "", s)
    if s in _NO_VALUE:
        return 0.0
    if not s:
        return None
    if "," in s:
        # German style decimal comma: 269.819,00 or 3,5
        int_part, _, dec_part = s.partition(",")
        int_part = int_part.replace(".", "")
        s = f"{int_part}.{dec_part}"
    elif re.fullmatch(r"-?\d+(\.\d{3})+", s):
        # Thousands grouping without decimals: 249.750
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


def occurrence_keys(header: list[str]) -> dict[tuple[str, int], int]:
    """Map (header_name, occurrence_0based) -> column index."""
    seen: dict[str, int] = {}
    out: dict[tuple[str, int], int] = {}
    for idx, name in enumerate(header):
        key = "" if name is None else str(name)
        occ = seen.get(key, 0)
        out[(key, occ)] = idx
        seen[key] = occ + 1
    return out


@dataclass
class FileSource:
    """A parsed file. Rows are dicts of (header_name, occurrence) -> value."""

    headers: list[str]
    rows: list[dict[tuple[str, int], object]]
    format: str  # 'xlsx' | 'csv'

    def header_index(self, name: str, occurrence: int = 0) -> int | None:
        cnt = 0
        for idx, h in enumerate(self.headers):
            hk = "" if h is None else str(h)
            if hk == name:
                if cnt == occurrence:
                    return idx
                cnt += 1
        return None

    def raw_mapping(self, index: int) -> str:
        """Human-readable raw key for a column index (de-duplicates headers)."""
        hk = "" if self.headers[index] is None else str(self.headers[index])
        cnt = sum(1 for h in self.headers[: index + 1] if ("" if h is None else str(h)) == hk)
        return hk if cnt == 1 else f"{hk} ({cnt})"


def _read_xlsx(stream) -> FileSource:
    wb = load_workbook(stream, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    try:
        header = [("" if h is None else str(h)) for h in next(it)]
    except StopIteration:
        wb.close()
        return FileSource([], [], "xlsx")
    occ = occurrence_keys(header)
    rows: list[dict[tuple[str, int], object]] = []
    for row in it:
        d: dict[tuple[str, int], object] = {}
        for (name, o), idx in occ.items():
            d[(name, o)] = row[idx] if idx < len(row) else None
        rows.append(d)
    wb.close()
    return FileSource(header, rows, "xlsx")


def _read_csv_bytes(data: bytes) -> FileSource:
    encodings = ("utf-8-sig", "utf-8", "cp1252")
    text = None
    for enc in encodings:
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Unsupported CSV encoding")
    sample = text[:4000]
    delimiter = ";"
    if "," in sample and ";" not in sample:
        delimiter = ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    try:
        header = ["".join(h.split("\ufeff")) if h else "" for h in next(reader)]
    except StopIteration:
        return FileSource([], [], "csv")
    header = [h.replace("\ufeff", "").strip() for h in header]
    occ = occurrence_keys(header)
    rows: list[dict[tuple[str, int], object]] = []
    for row in reader:
        d: dict[tuple[str, int], object] = {}
        for (name, o), idx in occ.items():
            d[(name, o)] = row[idx].strip() if idx < len(row) else ""
        rows.append(d)
    return FileSource(header, rows, "csv")


def read_file(path: Path | str) -> FileSource:
    p = Path(path)
    data = p.read_bytes()
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_xlsx(io.BytesIO(data))
    if p.suffix.lower() == ".csv":
        return _read_csv_bytes(data)
    raise ValueError(f"Unsupported file type: {p.suffix}")


def read_headers(path: Path | str) -> list[str]:
    """Read only the header row of a file (low memory for large xlsx)."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(p, read_only=True, data_only=True)
        try:
            ws = wb.worksheets[0]
            first = next(ws.iter_rows(values_only=True), None)
        finally:
            wb.close()
        if not first:
            return []
        return [("" if h is None else str(h)) for h in first]
    if p.suffix.lower() == ".csv":
        return _read_csv_bytes(p.read_bytes()).headers
    raise ValueError(f"Unsupported file type: {p.suffix}")


def rows_iterator(path: str | Path) -> Iterator[dict[tuple[str, int], object]]:
    """Streaming parser for large files (yields row dicts, low memory)."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = [("" if h is None else str(h)) for h in next(it)]
        occ = occurrence_keys(header)
        for row in it:
            d = {}
            for (name, o), idx in occ.items():
                d[(name, o)] = row[idx] if idx < len(row) else None
            yield d
        wb.close()
        return
    # CSV
    data = p.read_bytes()
    src = _read_csv_bytes(data)
    yield from src.rows
