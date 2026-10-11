import csv
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def is_empty(v):
    return v is None or str(v).strip() == ""


def check(name, key_col, path, fmt, req_cols):
    print(f"\n### {name}  key={key_col!r}  ({fmt})")
    vals = Counter()
    missing = Counter()
    n = 0

    if fmt == "xlsx":
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = list(next(it))
        d = dict(zip(header, range(len(header)), strict=False))
        for row in it:
            n += 1
            k = row[d.get(key_col)] if d.get(key_col) is not None else None
            vals[str(k) if k is not None else None] += 1
            for req in req_cols:
                idx = d.get(req)
                if idx is None or is_empty(row[idx]):
                    missing[req] += 1
        wb.close()
    else:
        with open(path, encoding="utf-8", errors="replace", newline="") as f:
            r = csv.reader(f, delimiter=";")
            header = next(r)
            d = dict(zip(header, range(len(header)), strict=False))
            for row in r:
                n += 1
                k = row[d.get(key_col)] if d.get(key_col) is not None else None
                vals[str(k) if k is not None else None] += 1
                for req in req_cols:
                    idx = d.get(req)
                    if idx is None or idx >= len(row) or is_empty(row[idx]):
                        missing[req] += 1

    dupes = {k: c for k, c in vals.items() if c > 1 and k is not None}
    print(f"rows={n} unique_keys={len(vals)} null_keys={vals.get(None, 0)} dupe_keys={len(dupes)}")
    if dupes:
        print("  DUPKEY sample:", sorted(dupes.items(), key=lambda x: -x[1])[:5])
    if missing:
        print("  MISSING:", dict(sorted(missing.items())))


B = Path("samples")
check(
    "Equipment",
    "Equipment",
    B / "equipment" / "equipment.xlsx",
    "xlsx",
    ["Equipment", "Functional Loc.", "Description"],
)
check(
    "FunctionalLocation",
    "Functional Loc.",
    B / "functional_location" / "functional_location.xlsx",
    "xlsx",
    ["Functional Loc.", "Description"],
)
check(
    "MeasuringDocument",
    "MeasDocument",
    B / "measuring_document" / "measuring_document.xlsx",
    "xlsx",
    ["MeasDocument", "Measuring point", "Date to be recorded"],
)
check(
    "MeasuringPoint",
    "Measuring point",
    B / "measuring_point" / "measuring_point.xlsx",
    "xlsx",
    ["Measuring point", "Equipment", "Characteristic"],
)
check(
    "Notification",
    "Notification",
    B / "notification" / "notification.csv",
    "csv",
    ["Notification", "Equipment", "Order"],
)
check(
    "WorkOrder",
    "Order",
    B / "work_order" / "work_order.csv",
    "csv",
    ["Order", "Notification", "Equipment"],
)
